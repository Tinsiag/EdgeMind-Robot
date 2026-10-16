from pathlib import Path
import hashlib
import json
import re
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[3]
project = root / 'code/fpga/acg720-motor-test'
archive = project / 'verification/2026-10-15'
vendor = Path('C:/Users/sc/Documents/Dev/Hardware/fpga/小梅哥ACG720国产高云FPGA教学开发板资料')
sheet = vendor / '04_引脚功能信息/高云 ACG720-138&60K通用管脚分配表.xlsx'
pdf = vendor / '06_硬件图纸/ACG720-138&60K通用原理图.pdf'
power_pdf = Path('C:/Users/sc/Downloads/2/SCH_Schematic1_2026-10-09.pdf')
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest().upper()
rows = [
    ('left_in1', 'LAIN1', 'L-1', 5, 'A18', 'out'),
    ('left_in2', 'LAIN2', 'L-2', 6, 'A19', 'out'),
    ('right_in1', 'RAIN1', 'R-5', 7, 'F13', 'out'),
    ('right_in2', 'RAIN2', 'R-6', 8, 'F14', 'out'),
    ('left_enc_a', 'LF_ENA', 'L-7', 3, 'A15', 'in'),
    ('left_enc_b', 'LF_ENB', 'L-8', 4, 'A16', 'in'),
    ('right_enc_a', 'RF_ENA', 'R-2', 9, 'E13', 'in'),
    ('right_enc_b', 'RF_ENB', 'R-1', 10, 'E14', 'in'),
]
wb = load_workbook(sheet, read_only=True, data_only=True)
vendor_gpio = {}
for row in wb.worksheets[0].iter_rows(values_only=True):
    values = [str(v) for v in row if v is not None]
    for i, value in enumerate(values[:-1]):
        if re.fullmatch(r'GPIO0_\d+', value):
            vendor_gpio[value] = values[i+1]
cst_path = project / 'constraints/acg720_motor_test.cst'
locs = dict(re.findall(r'IO_LOC\s+"([^"]+)"\s+(\w+);', cst_path.read_text(encoding='utf-8')))
report_path = project / 'impl/pnr/acg720_bridge_probe.rpt.txt'
report = report_path.read_text(encoding='utf-8')
section = report.split('7. Pinout by Port Name', 1)[1].split('8. All Package Pins', 1)[0]
actual = {}
for line in section.splitlines():
    cells = [s.strip() for s in line.split('|')]
    if len(cells) >= 18 and cells[0] in locs:
        actual[cells[0]] = (cells[2].split('/')[0], cells[4])
verified = []
for port, net, header, pin, ball, direction in rows:
    gpio = f'GPIO0_{pin-1}'  # P7 first ten pins, visually verified on vendor PDF page 11.
    assert vendor_gpio[gpio] == ball, (port, 'vendor spreadsheet')
    assert locs[port] == ball, (port, 'CST')
    assert actual[port] == (ball, direction), (port, 'actual routed output')
    verified.append(dict(port=port, net=net, carrier_header=header, P7_pin=pin,
                         column=(pin+1)//2, row='下' if pin%2 else '上',
                         ball=ball, direction=direction))
vg_path = project / 'impl/gwsynthesis/acg720_bridge_probe.vg'
vg = vg_path.read_text(encoding='utf-8')
buffers = {out.strip(): inp.strip() for out, inp in re.findall(
    r'OBUF\s+\w+\s*\(\s*\.O\(([^)]+)\),\s*\.I\(([^)]+)\)', vg)}
assert buffers['right_in1'] == buffers['led[2]'] == 'led_d[2]'
assert buffers['left_in1'] == buffers['led[1]'] == 'led_d[1]'
assert buffers['right_in2'] == buffers['left_in2'] == 'GND'
audit = json.loads((archive/'build-audit-bridge_probe-wiring-20261014.json').read_text(encoding='utf-8'))
assert sha(project/'impl/pnr/acg720_bridge_probe.fs') == audit['bitstream_sha256']
assert sha(report_path) == sha(archive/'acg720_bridge_probe.rpt.txt')
for name, expected in audit['source_sha256'].items():
    assert sha(project/name) == expected, name
assert len(audit['pins']) == 21 and audit['timing_violations'] == {'Setup': 0, 'Hold': 0}
result = {
    'record_date': '2026-10-15',
    'conclusion': 'No pin assignment, direction, channel selection or missing-enable error found in reviewed A71A software artifacts.',
    'verified_signal_paths': verified,
    'synthesized_output_buffers': {p: buffers[p] for p in ['left_in1','left_in2','right_in1','right_in2','led[1]','led[2]']},
    'source_and_bitstream_hashes_match_reviewed_build': True,
    'bitstream_sha256': audit['bitstream_sha256'],
    'vendor_pdf_page_11_and_power_schematic_visually_reviewed': True,
    'carrier_A_channel_paths_from_power_schematic': {
        'L-1': 'LAIN1 -> U3 pin 3 (AIN1)', 'L-2': 'LAIN2 -> U3 pin 4 (AIN2)',
        'R-5': 'RAIN1 -> U6 pin 3 (AIN1)', 'R-6': 'RAIN2 -> U6 pin 4 (AIN2)',
        'left_motor_power': 'U2 pins 2/1 (LAOUT1/2) -> left six-pin socket pins 6/1',
        'right_motor_power': 'U4 pins 2/1 (RAOUT1/2) -> right six-pin socket pins 6/1',
    },
    'physical_board_matches_pdf_proven': False,
    'physical_pad_voltages_proven_by_software_audit': False,
    'current_turn_user_code_read_attempts': {'default_2_5MHz_exit': 50, 'retry_1_5MHz_exit': 50, 'successful': False},
    'last_successful_user_code_read': '0x0000A71A (previous step)',
    'new_source_edit_build_download_reset_or_motion': False,
    'new_regression_results': {'A71A_simulations_passed': 2, 'static_pinmap_tests_passed': 10},
    'inputs_sha256': {str(p): sha(p) for p in [sheet,pdf,power_pdf,cst_path,vg_path,report_path]},
}
out = archive / 'confirmed-wiring-code-review-20261015.json'
out.write_text(json.dumps(result,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
print('PASS: 8 signal paths agree across vendor spreadsheet, CST and routed report.')
print('PASS: D1/left_in1 and D2/right_in1 share synthesized signal; IN2 buffers tie to GND.')
print('PASS: all A71A build-source hashes, routed report and bitstream match archived review.')
print('Hardware limit: fresh user-code read failed twice; no physical voltage or continuity claim.')
